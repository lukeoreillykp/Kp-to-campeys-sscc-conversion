import base64
import io
import json
import re
import urllib.parse
import uuid
from datetime import datetime, timedelta, date, time
import smtplib

from email.mime.text import MIMEText
from email.mime.application import MIMEApplication
from email.mime.multipart import MIMEMultipart

import pandas as pd
import pytz
import requests
import streamlit as st

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

============================================================

PAGE CONFIG

============================================================

st.set_page_config(
page_title="KP to Campeys SSCC Sender",
page_icon="📦",
layout="wide",
initial_sidebar_state="collapsed",
)

============================================================

CONSTANTS

============================================================

EXPECTED_COLUMNS = [
"Load Ref",
"SSCC",
"Item Code",
"Item Description",
"Quantity",
]

FINAL_COLUMNS = [
"Load Ref",
"Date",
"SSCC",
"Item Code",
"Item Description",
"Quantity",
"Movement/JDE Order Ref",
]

GITHUB_OWNER = "lukeoreillykp"
GITHUB_REPO = "Kp-to-campeys-sscc-conversion"
GITHUB_BRANCH = "main"
GITHUB_FOLDER = "saved_loads"

SKU_HISTORY_PATH = "sku_counts_history.csv"
CONTACT_LIST_PATH = "campeys contact list.txt"
PLANNER_PATH = "load_planner_requests.json"

EMAIL_TO = "kpsnacks@campeys.co.uk"
LUKE_EMAIL = "luke.oreilly@kpsnacks.com"
GRAYSON_EMAIL = "grayson.swan@kpsnacks.com"

EMAIL_CC = [
LUKE_EMAIL,
GRAYSON_EMAIL,
]

GMAIL_ADDRESS = "kp.ponte.csv@gmail.com"

AUTOSTORE_URL = "https://autostore-live.snacks.local/app"

============================================================

SESSION STATE

============================================================

DEFAULT_SESSION_STATE = {
"page": "home",
"process_complete": False,
"csv_text": None,
"filename": None,
"sku_counts": None,
"load_ref": None,
"output_df": None,
"github_result": None,
"history_df": None,
"planner_week": None,
"planner_selected_slot_id": None,
}

for key, value in DEFAULT_SESSION_STATE.items():
if key not in st.session_state:
st.session_state[key] = value

============================================================

CSS

============================================================

st.markdown(
"""
<style>
.block-container {
padding-top: 1.5rem;
padding-bottom: 2rem;
}

.stButton > button {  
    border-radius: 8px;  
    font-weight: 600;  
}  

.planner-day-header {  
    text-align: center;  
    font-weight: 700;  
    padding: 8px 3px;  
    background-color: #eeeeee;  
    border-radius: 6px;  
    margin-bottom: 8px;  
}  

.planner-empty {  
    min-height: 72px;  
}  

[class*="st-key-planner_cell_pending_"] button,  
[class*="st-key-planner_cell_assigned_"] button,  
[class*="st-key-planner_cell_cancelled_"] button {  
    min-height: 76px;  
    height: 76px;  
    width: 100%;  
    font-size: 1.25rem;  
    font-weight: 800;  
    border-radius: 8px;  
    margin-bottom: 4px;  
}  

[class*="st-key-planner_cell_pending_"] button {  
    background: #f2f2f2;  
    color: #222;  
    border: 1px solid #cfcfcf;  
}  

[class*="st-key-planner_cell_assigned_"] button {  
    background: #d9ead3;  
    color: #274e13;  
    border: 2px solid #70ad47;  
}  

[class*="st-key-planner_cell_cancelled_"] button {  
    background: #f4cccc;  
    color: #990000;  
    border: 2px solid #cc0000;  
}  

[class*="st-key-planner_cell_pending_"] button:hover,  
[class*="st-key-planner_cell_assigned_"] button:hover,  
[class*="st-key-planner_cell_cancelled_"] button:hover {  
    filter: brightness(0.97);  
}  

.planner-selected-title {  
    font-size: 1.05rem;  
    font-weight: 700;  
    margin-bottom: 0.35rem;  
}  

div[data-testid="stExpander"] {  
    border-radius: 8px;  
}  

.st-key-history_back button,  
.st-key-import_back button,  
.st-key-contacts_back button,  
.st-key-sender_back button,  
.st-key-planner_back button {  
    position: fixed;  
    top: 15px;  
    left: 15px;  
    z-index: 999999;  
    width: auto;  
    min-width: 120px;  
}  

@media (max-width: 800px) {  
    .st-key-history_back button,  
    .st-key-import_back button,  
    .st-key-contacts_back button,  
    .st-key-sender_back button,  
    .st-key-planner_back button {  
        position: static;  
        margin-bottom: 10px;  
    }  
}  
</style>  
""",  
unsafe_allow_html=True,

)

============================================================

NAVIGATION

============================================================

def go_to(page):
st.session_state.page = page

if page != "sender":  
    st.session_state.process_complete = False  
    st.session_state.csv_text = None  
    st.session_state.filename = None  
    st.session_state.sku_counts = None  
    st.session_state.load_ref = None  
    st.session_state.output_df = None  
    st.session_state.github_result = None

============================================================

GITHUB

============================================================

def get_github_settings():
token = st.secrets.get("github_token", "")
username = st.secrets.get("github_username", GITHUB_OWNER)
repo = st.secrets.get("github_repo", GITHUB_REPO)

return token, username, repo

def get_github_connection():
token, username, repo = get_github_settings()

if not token:  
    st.error("GitHub token is not configured in Streamlit Secrets.")  
    return None  

return {  
    "token": token,  
    "username": username,  
    "repo": repo,  
}

def github_headers(token):
return {
"Authorization": f"Bearer {token}",
"Accept": "application/vnd.github+json",
}

def get_github_file(file_path):
connection = get_github_connection()

if not connection:  
    return None  

token = connection["token"]  
repo = connection["repo"]  

url = (  
    f"https://api.github.com/repos/"  
    f"{GITHUB_OWNER}/{repo}/contents/{urllib.parse.quote(file_path)}"  
)  

response = requests.get(  
    url,  
    headers=github_headers(token),  
    params={"ref": GITHUB_BRANCH},  
    timeout=30,  
)  

if response.status_code == 404:  
    return None  

response.raise_for_status()  

data = response.json()  

if "content" not in data:  
    return None  

content = base64.b64decode(data["content"]).decode("utf-8-sig")  

return {  
    "content": content,  
    "sha": data.get("sha"),  
    "name": data.get("name"),  
    "path": data.get("path"),  
}

def upload_github_file(file_path, file_content, commit_message):
connection = get_github_connection()

if not connection:  
    return False, "GitHub connection unavailable."  

token = connection["token"]  
repo = connection["repo"]  

url = (  
    f"https://api.github.com/repos/"  
    f"{GITHUB_OWNER}/{repo}/contents/{urllib.parse.quote(file_path)}"  
)  

existing = get_github_file(file_path)  

encoded = base64.b64encode(  
    file_content.encode("utf-8")  
).decode("utf-8")  

payload = {  
    "message": commit_message,  
    "content": encoded,  
    "branch": GITHUB_BRANCH,  
}  

if existing and existing.get("sha"):  
    payload["sha"] = existing["sha"]  

response = requests.put(  
    url,  
    headers=github_headers(token),  
    json=payload,  
    timeout=30,  
)  

if response.status_code not in (200, 201):  
    return False, response.text  

return True, response.json()

def list_saved_load_files():
connection = get_github_connection()

if not connection:  
    return []  

token = connection["token"]  
repo = connection["repo"]  

url = (  
    f"https://api.github.com/repos/"  
    f"{GITHUB_OWNER}/{repo}/contents/{GITHUB_FOLDER}"  
)  

response = requests.get(  
    url,  
    headers=github_headers(token),  
    params={"ref": GITHUB_BRANCH},  
    timeout=30,  
)  

if response.status_code == 404:  
    return []  

if response.status_code != 200:  
    return []  

data = response.json()  

if not isinstance(data, list):  
    return []  

return sorted(  
    [  
        item["name"]  
        for item in data  
        if item.get("type") == "file"  
        and item.get("name", "").lower().endswith(".csv")  
    ]  
)

def find_saved_load_file_for_load_ref(load_ref, saved_files=None):
if not load_ref:
return None

if saved_files is None:  
    saved_files = list_saved_load_files()  

clean_ref = str(load_ref).strip().lower()  

for filename in saved_files:  
    stem = filename.rsplit(".", 1)[0].lower()  

    if stem == clean_ref or stem.startswith(clean_ref + "_"):  
        return filename  

return None

def upload_to_github(csv_text, filename):
file_path = f"{GITHUB_FOLDER}/{filename}"

return upload_github_file(  
    file_path,  
    csv_text,  
    f"Archive sender load {filename}",  
)

============================================================

HISTORY

============================================================

def load_history():
result = get_github_file(SKU_HISTORY_PATH)

if not result:  
    return pd.DataFrame(  
        columns=[  
            "Date",  
            "Load Ref",  
            "Item Code",  
            "Quantity",  
        ]  
    )  

try:  
    return pd.read_csv(  
        io.StringIO(result["content"]),  
        dtype=str,  
    )  
except Exception:  
    return pd.DataFrame()

def save_history(history_df):
csv_text = history_df.to_csv(index=False)

return upload_github_file(  
    SKU_HISTORY_PATH,  
    csv_text,  
    "Update SKU counts history",  
)

def update_sku_history(load_ref, sku_counts):
history = load_history()

new_rows = []  

for item_code, quantity in sku_counts.items():  
    new_rows.append(  
        {  
            "Date": datetime.now(  
                pytz.timezone("Europe/London")  
            ).strftime("%Y-%m-%d %H:%M:%S"),  
            "Load Ref": load_ref,  
            "Item Code": item_code,  
            "Quantity": int(quantity),  
        }  
    )  

if new_rows:  
    history = pd.concat(  
        [  
            history,  
            pd.DataFrame(new_rows),  
        ],  
        ignore_index=True,  
    )  

success, message = save_history(history)  

return success, message

============================================================

SENDER HELPERS

============================================================

def clean_column_names(df):
df = df.copy()

df.columns = [  
    str(column).strip()  
    for column in df.columns  
]  

return df

def find_column(df, possible_names):
normalised = {
re.sub(r"\s+", " ", str(column).strip().lower()): column
for column in df.columns
}

for name in possible_names:  
    key = re.sub(  
        r"\s+",  
        " ",  
        str(name).strip().lower(),  
    )  

    if key in normalised:  
        return normalised[key]  

return None

def read_wms_file(uploaded_file):
raw = uploaded_file.getvalue()

try:  
    df = pd.read_csv(  
        io.BytesIO(raw),  
        sep="\t",  
        dtype=str,  
        encoding="utf-8-sig",  
    )  

    if len(df.columns) <= 1:  
        raise ValueError()  

except Exception:  
    df = pd.read_csv(  
        io.BytesIO(raw),  
        dtype=str,  
        encoding="utf-8-sig",  
    )  

return clean_column_names(df)

def make_clean_load_ref(value):
if value is None or pd.isna(value):
return ""

text = str(value).strip()  

if text.endswith(".0"):  
    text = text[:-2]  

return text

def make_clean_sscc(value):
if value is None or pd.isna(value):
return ""

text = str(value).strip()  

if text.endswith(".0"):  
    text = text[:-2]  

return text

def process_wms_dataframe(df):
df = clean_column_names(df)

missing = [  
    column  
    for column in EXPECTED_COLUMNS  
    if column not in df.columns  
]  

if missing:  
    raise ValueError(  
        "Missing required columns: "  
        + ", ".join(missing)  
    )  

df = df.copy()  

for column in EXPECTED_COLUMNS:  
    df[column] = df[column].fillna("").astype(str).str.strip()  

df["Load Ref"] = df["Load Ref"].apply(make_clean_load_ref)  
df["SSCC"] = df["SSCC"].apply(make_clean_sscc)  

df = df[  
    (df["SSCC"] != "")  
    & (df["Item Code"] != "")  
].copy()  

summary_pattern = (  
    r"^\s*(total|totals|summary|grand total|"  
    r"sub[- ]?total)\s*$"  
)  

for column in ["SSCC", "Item Code", "Item Description"]:  
    if column in df.columns:  
        mask = ~df[column].str.contains(  
            summary_pattern,  
            case=False,  
            regex=True,  
            na=False,  
        )  
        df = df[mask].copy()  

if df.empty:  
    raise ValueError(  
        "No valid SSCC / Item Code rows were found."  
    )  

load_refs = [  
    ref for ref in df["Load Ref"].tolist()  
    if ref  
]  

if not load_refs:  
    raise ValueError(  
        "No Load Ref was found in the WMS file."  
    )  

load_ref = load_refs[0]  

london = pytz.timezone("Europe/London")  
now_london = datetime.now(london)  

df["Load Ref"] = load_ref  
df["Date"] = now_london.strftime(  
    "%d/%m/%Y %H:%M"  
)  

movement_col = find_column(  
    df,  
    [  
        "Movement/JDE Order Ref",  
        "Movement / JDE Order Ref",  
        "JDE Order Ref",  
        "Movement",  
        "Order Ref",  
    ],  
)  

if movement_col:  
    df["Movement/JDE Order Ref"] = df[  
        movement_col  
    ]  
else:  
    df["Movement/JDE Order Ref"] = ""  

df = df[  
    [  
        "Load Ref",  
        "Date",  
        "SSCC",  
        "Item Code",  
        "Item Description",  
        "Quantity",  
        "Movement/JDE Order Ref",  
    ]  
]  

return df, load_ref

def send_email_with_attachment(
csv_bytes,
filename,
load_ref,
own_email=None,
):
password = st.secrets.get(
"gmail_password",
"",
)

if not password:  
    return False, "Gmail password is not configured."  

recipients = list(EMAIL_CC)  

if own_email:  
    own_email = own_email.strip()  

    if own_email and own_email not in recipients:  
        recipients.append(own_email)  

message = MIMEMultipart()  

message["From"] = GMAIL_ADDRESS  
message["To"] = EMAIL_TO  
message["Cc"] = ", ".join(recipients)  
message["Subject"] = (  
    f"Campeys SSCC Load {load_ref}"  
)  

body = f"""

Hi,

Please find attached the SSCC file for load {load_ref}.

The load has been processed and archived by the KP to Campeys SSCC Sender.

Regards,

KP Pontefract
"""

message.attach(MIMEText(body, "plain"))  

attachment = MIMEApplication(  
    csv_bytes,  
    _subtype="csv",  
)  

attachment.add_header(  
    "Content-Disposition",  
    "attachment",  
    filename=filename,  
)  

message.attach(attachment)  

try:  
    with smtplib.SMTP_SSL(  
        "smtp.gmail.com",  
        465,  
        timeout=30,  
    ) as server:  
        server.login(  
            GMAIL_ADDRESS,  
            password,  
        )  

        server.sendmail(  
            GMAIL_ADDRESS,  
            [EMAIL_TO] + recipients,  
            message.as_string(),  
        )  

    return True, "Email sent successfully."  

except Exception as exc:  
    return False, str(exc)

============================================================

SENDER PAGE

============================================================

def show_sender():
if st.button(
"← Back to Home",
key="sender_back",
):
go_to("home")
st.rerun()

st.title("📦 KP to Campeys SSCC Sender")  

st.write(  
    "Upload the WMS export, clean it, archive it to GitHub "  
    "and optionally send it to Campeys."  
)  

uploaded_file = st.file_uploader(  
    "Upload WMS file",  
    type=["csv", "txt"],  
    key="wms_upload",  
)  

if uploaded_file is None:  
    return  

if st.button(  
    "Process WMS File",  
    type="primary",  
    use_container_width=True,  
):  
    try:  
        with st.spinner("Processing WMS file..."):  
            df = read_wms_file(uploaded_file)  

            output_df, load_ref = process_wms_dataframe(df)  

            sku_counts = (  
                output_df  
                .groupby("Item Code")["Quantity"]  
                .sum()  
                .to_dict()  
            )  

            london = pytz.timezone("Europe/London")  

            filename = (  
                f"{load_ref}_"  
                f"{datetime.now(london).strftime('%Y%m%d_%H%M%S')}.csv"  
            )  

            csv_text = output_df.to_csv(  
                index=False,  
                encoding="utf-8-sig",  
            )  

            github_ok, github_result = upload_to_github(  
                csv_text,  
                filename,  
            )  

            if github_ok:  
                history_ok, history_result = (  
                    update_sku_history(  
                        load_ref,  
                        sku_counts,  
                    )  
                )  
            else:  
                history_ok = False  
                history_result = (  
                    "History not updated because "  
                    "the GitHub archive failed."  
                )  

            st.session_state.process_complete = True  
            st.session_state.csv_text = csv_text  
            st.session_state.filename = filename  
            st.session_state.sku_counts = sku_counts  
            st.session_state.load_ref = load_ref  
            st.session_state.output_df = output_df  
            st.session_state.github_result = (  
                github_ok,  
                github_result,  
            )  

            if github_ok:  
                st.success(  
                    f"Load {load_ref} archived successfully."  
                )  
            else:  
                st.error(  
                    "The load could not be archived to GitHub."  
                )  

            if history_ok:  
                st.success(  
                    "SKU history updated."  
                )  
            else:  
                st.warning(  
                    f"SKU history was not updated: "  
                    f"{history_result}"  
                )  

    except Exception as exc:  
        st.error(  
            f"Unable to process the file: {exc}"  
        )  
        return  

if not st.session_state.process_complete:  
    return  

output_df = st.session_state.output_df  
csv_text = st.session_state.csv_text  
filename = st.session_state.filename  
load_ref = st.session_state.load_ref  

st.divider()  

st.subheader(  
    f"Processed Load: {load_ref}"  
)  

st.dataframe(  
    output_df,  
    use_container_width=True,  
    hide_index=True,  
)  

st.subheader("SKU Counts")  

sku_df = pd.DataFrame(  
    [  
        {  
            "Item Code": item,  
            "Quantity": quantity,  
        }  
        for item, quantity in (  
            st.session_state.sku_counts or {}  
        ).items()  
    ]  
)  

st.dataframe(  
    sku_df,  
    use_container_width=True,  
    hide_index=True,  
)  

st.download_button(  
    "⬇️ Download CSV",  
    data=csv_text.encode("utf-8-sig"),  
    file_name=filename,  
    mime="text/csv",  
    use_container_width=True,  
)  

st.subheader("Send to Campeys")  

own_email = st.text_input(  
    "Optional additional email address",  
    placeholder="your.email@example.com",  
)  

if st.button(  
    "📧 Send CSV to Campeys",  
    type="primary",  
    use_container_width=True,  
):  
    with st.spinner("Sending email..."):  
        success, message = send_email_with_attachment(  
            csv_text.encode("utf-8-sig"),  
            filename,  
            load_ref,  
            own_email,  
        )  

    if success:  
        st.success(message)  
    else:  
        st.error(  
            f"Email failed: {message}"  
        )

============================================================

PLANNER DATE/TIME HELPERS

============================================================

def normalise_date_only(value):
if value is None:
return None

if isinstance(value, datetime):  
    return value.date()  

if isinstance(value, date):  
    return value  

parsed = pd.to_datetime(  
    value,  
    errors="coerce",  
    dayfirst=True,  
)  

if pd.isna(parsed):  
    return None  

return parsed.date()

def parse_time_value(value):
if value is None:
return None

try:  
    if pd.isna(value):  
        return None  
except Exception:  
    pass  

if isinstance(value, time):  
    return value.strftime("%H:%M")  

if isinstance(value, (datetime, pd.Timestamp)):  
    return value.strftime("%H:%M")  

if isinstance(value, (int, float)):  
    number = float(value)  

    if 0 <= number < 1:  
        total_minutes = round(  
            number * 24 * 60  
        )  

        hours = (  
            total_minutes // 60  
        ) % 24  

        minutes = (  
            total_minutes % 60  
        )  

        return (  
            f"{hours:02d}:"  
            f"{minutes:02d}"  
        )  

text = str(value).strip()  

if not text:  
    return None  

match = re.fullmatch(  
    r"(\d{1,2}):(\d{2})(?::(\d{2}))?",  
    text,  
)  

if match:  
    hours = int(match.group(1))  
    minutes = int(match.group(2))  

    if (  
        0 <= hours <= 23  
        and 0 <= minutes <= 59  
    ):  
        return (  
            f"{hours:02d}:"  
            f"{minutes:02d}"  
        )  

parsed = pd.to_datetime(  
    text,  
    errors="coerce",  
)  

if pd.notna(parsed):  
    return parsed.strftime("%H:%M")  

match = re.search(  
    r"\b(\d{1,2}):(\d{2})\s*([APap][Mm])?\b",  
    text,  
)  

if match:  
    hours = int(match.group(1))  
    minutes = int(match.group(2))  
    meridian = match.group(3)  

    if meridian:  
        meridian = meridian.upper()  

        if hours == 12:  
            hours = 0  

        if meridian == "PM":  
            hours += 12  

    if (  
        0 <= hours <= 23  
        and 0 <= minutes <= 59  
    ):  
        return (  
            f"{hours:02d}:"  
            f"{minutes:02d}"  
        )  

return None

def time_sort_value(time_text):
try:
hours, minutes = (
str(time_text).split(":")
)

return int(hours) * 60 + int(minutes)  

except Exception:  
    return 9999

def monday_for(value):
if isinstance(value, datetime):
value = value.date()

if not isinstance(value, date):  
    value = normalise_date_only(value)  

if value is None:  
    value = date.today()  

return value - timedelta(  
    days=value.weekday()  
)

def format_week_label(monday):
sunday = monday + timedelta(days=6)

return (  
    f"{monday.strftime('%d/%m/%Y')}"  
    f" - "  
    f"{sunday.strftime('%d/%m/%Y')}"  
)

============================================================

PLANNER GITHUB DATA

============================================================

def empty_planner_data():
return {
"version": 1,
"updated_at": "",
"slots": [],
}

def load_planner_data():
result = get_github_file(
PLANNER_PATH
)

if not result:  
    return empty_planner_data()  

try:  
    data = json.loads(  
        result["content"]  
    )  

    if not isinstance(data, dict):  
        return empty_planner_data()  

    if "slots" not in data:  
        data["slots"] = []  

    return data  

except Exception:  
    return empty_planner_data()

def save_planner_data(planner_data):
planner_data = dict(
planner_data
)

london = pytz.timezone(  
    "Europe/London"  
)  

planner_data["updated_at"] = (  
    datetime.now(london).isoformat()  
)  

content = json.dumps(  
    planner_data,  
    indent=2,  
    ensure_ascii=False,  
)  

return upload_github_file(  
    PLANNER_PATH,  
    content,  
    "Update Campeys load planner",  
)

============================================================

ARCHIVED LOADS

============================================================

def load_archived_loads():
files = list_saved_load_files()

archived = []  

for filename in files:  
    path = (  
        f"{GITHUB_FOLDER}/"  
        f"{filename}"  
    )  

    result = get_github_file(path)  

    if not result:  
        continue  

    try:  
        df = pd.read_csv(  
            io.StringIO(  
                result["content"]  
            ),  
            dtype=str,  
        )  

        if df.empty:  
            continue  

        load_ref = None  

        if "Load Ref" in df.columns:  
            for value in df["Load Ref"]:  
                if (  
                    value is not None  
                    and not pd.isna(value)  
                    and str(value).strip()  
                ):  
                    load_ref = str(  
                        value  
                    ).strip()  
                    break  

        if not load_ref:  
            continue  

        date_value = None  

        for column in [  
            "Date",  
            "Date Submitted",  
            "date",  
        ]:  
            if column in df.columns:  
                for value in df[column]:  
                    parsed = pd.to_datetime(  
                        value,  
                        errors="coerce",  
                        dayfirst=True,  
                    )  

                    if pd.notna(parsed):  
                        date_value = parsed  
                        break  

            if date_value is not None:  
                break  

        if date_value is None:  
            continue  

        archived.append(  
            {  
                "load_ref": load_ref,  
                "date": date_value.date().isoformat(),  
                "datetime": date_value,  
                "filename": filename,  
                "path": path,  
            }  
        )  

    except Exception:  
        continue  

archived.sort(  
    key=lambda item: (  
        item["date"],  
        item["datetime"],  
        item["filename"],  
    )  
)  

return archived

============================================================

PLANNER ARCHIVE MATCHING

============================================================

def sync_planner_with_archived_loads(
planner_data,
archived_loads,
):
slots = planner_data.get(
"slots",
[],
)

changed = False  

used_load_refs = {  
    str(slot.get("load_ref")).strip()  
    for slot in slots  
    if slot.get("load_ref")  
    and slot.get("status") in (  
        "assigned",  
        "cancelled",  
    )  
}  

archived_by_date = {}  

for load in archived_loads:  
    archived_by_date.setdefault(  
        load["date"],  
        [],  
    ).append(load)  

for date_key, loads in archived_by_date.items():  
    day_slots = [  
        slot  
        for slot in slots  
        if slot.get("collection_date") == date_key  
    ]  

    day_slots.sort(  
        key=lambda slot: (  
            time_sort_value(  
                slot.get(  
                    "collection_time",  
                    "23:59",  
                )  
            ),  
            slot.get("id", ""),  
        )  
    )  

    loads = sorted(  
        loads,  
        key=lambda item: (  
            item["datetime"],  
            item["filename"],  
        ),  
    )  

    for load in loads:  
        load_ref = load["load_ref"]  

        if load_ref in used_load_refs:  
            continue  

        pending_slot = next(  
            (  
                slot  
                for slot in day_slots  
                if slot.get("status") == "pending"  
                and not slot.get("load_ref")  
            ),  
            None,  
        )  

        if pending_slot is None:  
            break  

        pending_slot["status"] = "assigned"  
        pending_slot["load_ref"] = load_ref  
        pending_slot["archive_filename"] = (  
            load["filename"]  
        )  
        pending_slot["assigned_at"] = (  
            datetime.now(  
                pytz.timezone(  
                    "Europe/London"  
                )  
            ).isoformat()  
        )  

        used_load_refs.add(load_ref)  
        changed = True  

return changed

============================================================

REQUEST WORKBOOK IMPORTER

============================================================

def find_header_row(raw_df):
for row_index in range(
min(len(raw_df), 20)
):
values = [
str(value).strip().lower()
if value is not None
else ""
for value in raw_df.iloc[
row_index
].tolist()
]

has_date = any(  
        value == "date"  
        for value in values  
    )  

    has_planned_time = any(  
        value in (  
            "planned time",  
            "planned_time",  
        )  
        for value in values  
    )  

    if has_date and has_planned_time:  
        return row_index  

return None

def parse_campeys_request_sheet(
raw_df,
sheet_name,
):
header_row = find_header_row(
raw_df
)

if header_row is None:  
    return []  

headers = [  
    str(value).strip()  
    if value is not None  
    else ""  
    for value in raw_df.iloc[  
        header_row  
    ].tolist()  
]  

data = raw_df.iloc[  
    header_row + 1:  
].copy()  

data.columns = headers  

data = data.loc[  
    :,  
    [  
        column  
        for column in data.columns  
        if column  
    ],  
]  

date_col = find_column(  
    data,  
    ["Date"],  
)  

site_col = find_column(  
    data,  
    ["Del Site"],  
)  

planned_time_col = find_column(  
    data,  
    [  
        "Planned Time",  
        "Planned_Time",  
    ],  
)  

if not date_col or not planned_time_col:  
    return []  

imported = []  

for _, row in data.iterrows():  
    collection_date = normalise_date_only(  
        row.get(date_col)  
    )  

    collection_time = parse_time_value(  
        row.get(planned_time_col)  
    )  

    if (  
        collection_date is None  
        or collection_time is None  
    ):  
        continue  

    if site_col:  
        site = str(  
            row.get(site_col, "")  
        ).strip().lower()  

        if site and site != "campey":  
            continue  

    imported.append(  
        {  
            "collection_date": (  
                collection_date.isoformat()  
            ),  
            "collection_time": collection_time,  
            "source": (  
                f"Campey Transport Requests - "  
                f"{str(sheet_name).strip()}"  
            ),  
        }  
    )  

return imported

def parse_collection_requests_workbook(
uploaded_file,
):
imported = []

xls = pd.ExcelFile(  
    uploaded_file,  
    engine="openpyxl",  
)  

for sheet_name in xls.sheet_names:  
    try:  
        raw = pd.read_excel(  
            xls,  
            sheet_name=sheet_name,  
            header=None,  
            dtype=object,  
        )  
    except Exception:  
        continue  

    sheet_records = (  
        parse_campeys_request_sheet(  
            raw,  
            sheet_name,  
        )  
    )  

    imported.extend(  
        sheet_records  
    )  

return imported

def merge_planner_requests(
planner_data,
imported_requests,
):
slots = planner_data.setdefault(
"slots",
[],
)

added = 0  
existing = 0  

existing_keys = {}  

for slot in slots:  
    key = (  
        slot.get("collection_date"),  
        slot.get("collection_time"),  
    )  

    existing_keys.setdefault(  
        key,  
        [],  
    ).append(slot)  

imported_seen = {}  

for request in imported_requests:  
    date_key = request[  
        "collection_date"  
    ]  

    time_key = request[  
        "collection_time"  
    ]  

    key = (  
        date_key,  
        time_key,  
    )  

    imported_seen[key] = (  
        imported_seen.get(key, 0)  
        + 1  
    )  

    requested_number = (  
        imported_seen[key]  
    )  

    existing_for_key = (  
        existing_keys.get(key, [])  
    )  

    if requested_number <= len(  
        existing_for_key  
    ):  
        existing += 1  
        continue  

    slot = {  
        "id": str(  
            uuid.uuid4()  
        ),  
        "collection_date": date_key,  
        "collection_time": time_key,  
        "status": "pending",  
        "load_ref": "",  
        "archive_filename": "",  
        "notes": "",  
        "source": request.get(  
            "source",  
            "collection_request",  
        ),  
        "created_at": datetime.now(  
            pytz.timezone(  
                "Europe/London"  
            )  
        ).isoformat(),  
    }  

    slots.append(slot)  
    existing_for_key.append(slot)  
    existing_keys[key] = existing_for_key  

    added += 1  

return added, existing

============================================================

PLANNER EXPORT

============================================================

def create_planner_excel(
planner_data,
archived_loads,
):
wb = Workbook()

ws = wb.active  
ws.title = "Weekly Planner"  

headers = [  
    "Date",  
    "Day",  
    "Collection Time",  
    "Status",  
    "Display",  
    "Load Ref",  
    "Archive File",  
    "Notes",  
]  

ws.append(headers)  

header_fill = PatternFill(  
    fill_type="solid",  
    fgColor="D9EAD3",  
)  

green_fill = PatternFill(  
    fill_type="solid",  
    fgColor="D9EAD3",  
)  

red_fill = PatternFill(  
    fill_type="solid",  
    fgColor="F4CCCC",  
)  

grey_fill = PatternFill(  
    fill_type="solid",  
    fgColor="F2F2F2",  
)  

for cell in ws[1]:  
    cell.fill = header_fill  
    cell.font = Font(  
        bold=True  
    )  
    cell.alignment = Alignment(  
        horizontal="center"  
    )  

sorted_slots = sorted(  
    planner_data.get(  
        "slots",  
        [],  
    ),  
    key=lambda slot: (  
        slot.get(  
            "collection_date",  
            "",  
        ),  
        time_sort_value(  
            slot.get(  
                "collection_time",  
                "23:59",  
            )  
        ),  
    ),  
)  

for slot in sorted_slots:  
    status = slot.get(  
        "status",  
        "pending",  
    )  

    if status == "assigned":  
        display = slot.get(  
            "load_ref",  
            "",  
        )  
    elif status == "cancelled":  
        display = "C"  
    else:  
        display = "1"  

    collection_date = (  
        slot.get(  
            "collection_date",  
            "",  
        )  
    )  

    try:  
        day_name = datetime.strptime(  
            collection_date,  
            "%Y-%m-%d",  
        ).strftime("%A")  
    except Exception:  
        day_name = ""  

    ws.append(  
        [  
            collection_date,  
            day_name,  
            slot.get(  
                "collection_time",  
                "",  
            ),  
            status,  
            display,  
            slot.get(  
                "load_ref",  
                "",  
            ),  
            slot.get(  
                "archive_filename",  
                "",  
            ),  
            slot.get(  
                "notes",  
                "",  
            ),  
        ]  
    )  

    row_number = ws.max_row  

    if status == "assigned":  
        fill = green_fill  
    elif status == "cancelled":  
        fill = red_fill  
    else:  
        fill = grey_fill  

    ws.cell(  
        row=row_number,  
        column=5,  
    ).fill = fill  

for column in range(  
    1,  
    ws.max_column + 1,  
):  
    ws.column_dimensions[  
        get_column_letter(column)  
    ].width = 20  

archive_ws = wb.create_sheet(  
    "Archived Loads"  
)  

archive_headers = [  
    "Date",  
    "Load Ref",  
    "Archived CSV",  
]  

archive_ws.append(  
    archive_headers  
)  

for cell in archive_ws[1]:  
    cell.fill = header_fill  
    cell.font = Font(  
        bold=True  
    )  

for load in archived_loads:  
    archive_ws.append(  
        [  
            load.get(  
                "date",  
                "",  
            ),  
            load.get(  
                "load_ref",  
                "",  
            ),  
            load.get(  
                "filename",  
                "",  
            ),  
        ]  
    )  

for column in range(  
    1,  
    archive_ws.max_column + 1,  
):  
    archive_ws.column_dimensions[  
        get_column_letter(column)  
    ].width = 25  

output = io.BytesIO()  

wb.save(output)  

output.seek(0)  

return output.getvalue()

============================================================

PLANNER DISPLAY

============================================================

def planner_slot_display(slot):
status = slot.get(
"status",
"pending",
)

if status == "assigned":  
    return (  
        slot.get(  
            "load_ref"  
        )  
        or "1",  
        "planner-assigned",  
    )  

if status == "cancelled":  
    return (  
        "C",  
        "planner-cancelled",  
    )  

return (  
    "1",  
    "planner-pending",  
)

def change_slot_status(
planner_data,
slot_id,
action,
):
for slot in planner_data.get(
"slots",
[],
):
if slot.get("id") != slot_id:
continue

if action in (  
        "Cancel load",  
        "Cancel collection",  
    ):  
        slot["status"] = "cancelled"  

        slot["cancelled_at"] = (  
            datetime.now(  
                pytz.timezone(  
                    "Europe/London"  
                )  
            ).isoformat()  
        )  

        return True  

    if action in (  
        "Re-open",  
        "Re-open collection",  
    ):  
        slot["status"] = "pending"  

        slot["load_ref"] = ""  
        slot["archive_filename"] = ""  

        slot["reopened_at"] = (  
            datetime.now(  
                pytz.timezone(  
                    "Europe/London"  
                )  
            ).isoformat()  
        )  

        return True  

return False

============================================================

LOAD PLANNER

============================================================

def show_planner():
if st.button(
"← Back to Home",
key="planner_back",
):
go_to("home")
st.rerun()

st.title("📋 Campeys Load Planner")  

st.caption(  
    "Collection requests come from the existing Campeys "  
    "Transport Requests workbook. The planner uses "  
    "Date + Planned Time."  
)  

# --------------------------------------------------------  
# WEEK CONTROL  
# --------------------------------------------------------  

if st.session_state.planner_week is None:  
    st.session_state.planner_week = monday_for(  
        date.today()  
    )  

current_monday = monday_for(  
    st.session_state.planner_week  
)  

previous_col, week_col, today_col, next_col = st.columns(  
    [1.15, 2.2, 0.8, 1.15]  
)  

with previous_col:  
    if st.button(  
        "⬅ Previous Week",  
        use_container_width=True,  
    ):  
        st.session_state.planner_week = (  
            current_monday  
            - timedelta(days=7)  
        )  
        st.session_state.planner_selected_slot_id = None  
        st.rerun()  

with week_col:  
    st.markdown(  
        f"""  
        <div style="  
            text-align:center;  
            font-size:1.15rem;  
            font-weight:700;  
            padding-top:0.35rem;  
        ">  
            {format_week_label(current_monday)}  
        </div>  
        """,  
        unsafe_allow_html=True,  
    )  

with today_col:  
    if st.button(  
        "Today",  
        use_container_width=True,  
    ):  
        st.session_state.planner_week = (  
            monday_for(date.today())  
        )  
        st.session_state.planner_selected_slot_id = None  
        st.rerun()  

with next_col:  
    if st.button(  
        "Next Week ➡",  
        use_container_width=True,  
    ):  
        st.session_state.planner_week = (  
            current_monday  
            + timedelta(days=7)  
        )  
        st.session_state.planner_selected_slot_id = None  
        st.rerun()  

current_monday = monday_for(  
    st.session_state.planner_week  
)  

# --------------------------------------------------------  
# IMPORT REQUESTS  
# --------------------------------------------------------  

with st.expander(  
    "⚙️ Planner Settings / Import Requests",  
    expanded=False,  
):  
    st.write(  
        "Upload the existing Campeys Transport Requests "  
        "workbook. The planner reads **Date** and "  
        "**Planned Time**. The blank **Collection Time** "  
        "column is not used."  
    )  

    request_file = st.file_uploader(  
        "Campey Transport Requests workbook",  
        type=["xlsx", "xls"],  
        key="planner_request_upload",  
    )  

    if request_file is not None:  
        try:  
            requests_found = (  
                parse_collection_requests_workbook(  
                    request_file  
                )  
            )  

            if not requests_found:  
                st.warning(  
                    "No Campeys collection requests were "  
                    "found in the workbook."  
                )  
            else:  
                preview_df = pd.DataFrame(  
                    requests_found  
                )  

                preview_df = (  
                    preview_df  
                    .sort_values(  
                        [  
                            "collection_date",  
                            "collection_time",  
                        ]  
                    )  
                    .reset_index(drop=True)  
                )  

                st.success(  
                    f"Found {len(requests_found)} "  
                    f"collection request(s)."  
                )  

                st.dataframe(  
                    preview_df,  
                    use_container_width=True,  
                    hide_index=True,  
                )  

                if st.button(  
                    "Import / Update Planner",  
                    type="primary",  
                    key="import_planner_requests",  
                    use_container_width=True,  
                ):  
                    planner_data = (  
                        load_planner_data()  
                    )  

                    added, existing = (  
                        merge_planner_requests(  
                            planner_data,  
                            requests_found,  
                        )  
                    )  

                    archived_loads = (  
                        load_archived_loads()  
                    )  

                    sync_changed = (  
                        sync_planner_with_archived_loads(  
                            planner_data,  
                            archived_loads,  
                        )  
                    )  

                    success, message = (  
                        save_planner_data(  
                            planner_data  
                        )  
                    )  

                    if success:  
                        st.success(  
                            f"Planner updated. "  
                            f"{added} new slot(s) added; "  
                            f"{existing} existing slot(s) "  
                            f"kept."  
                        )  

                        if sync_changed:  
                            st.info(  
                                "Archived loads were also "  
                                "matched to pending "  
                                "collections."  
                            )  

                        st.rerun()  

                    else:  
                        st.error(  
                            f"Could not save planner: "  
                            f"{message}"  
                        )  

        except Exception as exc:  
            st.error(  
                f"Could not read the workbook: {exc}"  
            )  

# --------------------------------------------------------  
# REFRESH / ARCHIVE SYNC  
# --------------------------------------------------------  

refresh_col1, refresh_col2 = st.columns(  
    [1, 3]  
)  

with refresh_col1:  
    if st.button(  
        "🔄 Refresh from GitHub",  
        use_container_width=True,  
    ):  
        st.session_state.planner_selected_slot_id = None  
        st.rerun()  

planner_data = load_planner_data()  

archived_loads = load_archived_loads()  

sync_changed = (  
    sync_planner_with_archived_loads(  
        planner_data,  
        archived_loads,  
    )  
)  

if sync_changed:  
    success, message = (  
        save_planner_data(  
            planner_data  
        )  
    )  

    if success:  
        st.success(  
            "New archived load(s) matched to "  
            "the first pending collection(s)."  
        )  
    else:  
        st.warning(  
            "Loads were matched locally but the "  
            f"planner could not be saved: {message}"  
        )  

# --------------------------------------------------------  
# CURRENT WEEK  
# --------------------------------------------------------  

week_dates = [  
    current_monday  
    + timedelta(days=offset)  
    for offset in range(7)  
]  

week_date_keys = {  
    current_date.isoformat()  
    for current_date in week_dates  
}  

week_slots = [  
    slot  
    for slot in planner_data.get(  
        "slots",  
        [],  
    )  
    if slot.get(  
        "collection_date"  
    ) in week_date_keys  
]  

# Build all times in the week.  
times = sorted(  
    {  
        slot.get(  
            "collection_time",  
            "",  
        )  
        for slot in week_slots  
        if slot.get(  
            "collection_time"  
        )  
    },  
    key=time_sort_value,  
)  

# --------------------------------------------------------  
# WEEKLY PLANNER GRID  
# --------------------------------------------------------  

st.subheader("Weekly Planner")  

if not times:  
    st.info(  
        "No collection requests have been imported "  
        "for this week yet."  
    )  
else:  
    # Header.  
    header_cols = st.columns(  
        [0.7] + [1] * 7  
    )  

    header_cols[0].markdown(  
        "**Time**"  
    )  

    for index, current_date in enumerate(  
        week_dates,  
        start=1,  
    ):  
        day_key = current_date.isoformat()  

        count = sum(  
            1  
            for slot in week_slots  
            if slot.get(  
                "collection_date"  
            ) == day_key  
        )  

        header_cols[index].markdown(  
            f"""  
            <div class="planner-day-header">  
                {current_date.strftime('%A')}<br>  
                {current_date.strftime('%d/%m')}  
                <br>  
                <small>{count} collection(s)</small>  
            </div>  
            """,  
            unsafe_allow_html=True,  
        )  

    st.markdown(  
        "**Legend:** 🩶 Pending = 1 &nbsp;&nbsp; "  
        "🟩 Archived = Load Ref &nbsp;&nbsp; "  
        "🟥 Cancelled = C"  
    )  

    # Each row is a collection time.  
    for time_text in times:  
        row_cols = st.columns(  
            [0.7] + [1] * 7  
        )  

        row_cols[0].markdown(  
            f"**{time_text}**"  
        )  

        for day_index, current_date in enumerate(  
            week_dates,  
            start=1,  
        ):  
            date_key = (  
                current_date.isoformat()  
            )  

            matching_slots = [  
                slot  
                for slot in week_slots  
                if slot.get(  
                    "collection_date"  
                ) == date_key  
                and slot.get(  
                    "collection_time"  
                ) == time_text  
            ]  

            cell = row_cols[  
                day_index  
            ]  

            if not matching_slots:  
                cell.markdown(  
                    '<div class="planner-empty"></div>',  
                    unsafe_allow_html=True,  
                )  
                continue  

            for slot_number, slot in enumerate(  
                matching_slots  
            ):  
                display_value, _ = (  
                    planner_slot_display(  
                        slot  
                    )  
                )  

                status = slot.get(  
                    "status",  
                    "pending",  
                )  

                if status == "assigned":  
                    status_key = "assigned"  
                elif status == "cancelled":  
                    status_key = "cancelled"  
                else:  
                    status_key = "pending"  

                if cell.button(  
                    display_value,  
                    key=(  
                        f"planner_cell_"  
                        f"{status_key}_"  
                        f"{slot.get('id')}_"  
                        f"{slot_number}"  
                    ),  
                    use_container_width=True,  
                    help=(  
                        f"{current_date.strftime('%A %d/%m')}"  
                        f" — {time_text}"  
                    ),  
                ):  
                    st.session_state.planner_selected_slot_id = (  
                        slot.get("id")  
                    )  
                    st.rerun()  

    # ----------------------------------------------------  
    # SELECTED COLLECTION CONTROL  
    # ----------------------------------------------------  

    selected_slot_id = (  
        st.session_state.get(  
            "planner_selected_slot_id"  
        )  
    )  

    selected_slot = next(  
        (  
            slot  
            for slot in planner_data.get(  
                "slots",  
                [],  
            )  
            if slot.get("id")  
            == selected_slot_id  
        ),  
        None,  
    )  

    if selected_slot:  
        selected_date = normalise_date_only(  
            selected_slot.get(  
                "collection_date"  
            )  
        )  

        date_text = (  
            selected_date.strftime(  
                "%A %d/%m"  
            )  
            if selected_date  
            else selected_slot.get(  
                "collection_date",  
                "",  
            )  
        )  

        selected_time = selected_slot.get(  
            "collection_time",  
            "",  
        )  

        status = selected_slot.get(  
            "status",  
            "pending",  
        )  

        if status == "assigned":  
            status_label = "Archived"  
        elif status == "cancelled":  
            status_label = "Cancelled"  
        else:  
            status_label = "Pending"  

        st.markdown(  
            "#### Selected Collection"  
        )  

        with st.container(border=True):  
            info_col1, info_col2 = st.columns(  
                [1, 1]  
            )  

            with info_col1:  
                st.markdown(  
                    f"""  
                    <div class="planner-selected-title">  
                        {date_text} — {selected_time}  
                    </div>  
                    """,  
                    unsafe_allow_html=True,  
                )  

                if status == "assigned":  
                    st.markdown(  
                        f"**Load Ref:** "  
                        f"{selected_slot.get('load_ref', '')}"  
                    )  

                st.markdown(  
                    f"**Status:** {status_label}"  
                )  

            with info_col2:  
                if status == "cancelled":  
                    options = [  
                        "Re-open",  
                    ]  
                else:  
                    options = [  
                        "No change",  
                        "Cancel load",  
                    ]  

                action = st.selectbox(  
                    "Action",  
                    options,  
                    key=(  
                        "planner_selected_action_"  
                        f"{selected_slot.get('id')}"  
                    ),  
                )  

            if action != "No change":  
                changed = change_slot_status(  
                    planner_data,  
                    selected_slot.get("id"),  
                    action,  
                )  

                if changed:  
                    success, message = (  
                        save_planner_data(  
                            planner_data  
                        )  
                    )  

                    if not success:  
                        st.error(  
                            "Could not save planner "  
                            f"change: {message}"  
                        )  

                    st.session_state.planner_selected_slot_id = None  
                    st.rerun()  

# --------------------------------------------------------  
# ARCHIVED LOADS  
# --------------------------------------------------------  

with st.expander(  
    "📦 Archived Loads Detected",  
    expanded=False,  
):  
    current_week_archived = [  
        load  
        for load in archived_loads  
        if load.get(  
            "date"  
        ) in week_date_keys  
    ]  

    if not current_week_archived:  
        st.info(  
            "No archived sender CSVs were found "  
            "for this week."  
        )  
    else:  
        archive_rows = []  

        for load in current_week_archived:  
            assigned_slot = next(  
                (  
                    slot  
                    for slot in planner_data.get(  
                        "slots",  
                        [],  
                    )  
                    if slot.get(  
                        "load_ref"  
                    ) == load.get(  
                        "load_ref"  
                    )  
                ),  
                None,  
            )  

            archive_rows.append(  
                {  
                    "Date": load.get(  
                        "date",  
                        "",  
                    ),  
                    "Load Ref": load.get(  
                        "load_ref",  
                        "",  
                    ),  
                    "Archived CSV": load.get(  
                        "filename",  
                        "",  
                    ),  
                    "Planner Status": (  
                        assigned_slot.get(  
                            "status",  
                            "",  
                        )  
                        if assigned_slot  
                        else "Not assigned"  
                    ),  
                    "Collection Time": (  
                        assigned_slot.get(  
                            "collection_time",  
                            "",  
                        )  
                        if assigned_slot  
                        else ""  
                    ),  
                }  
            )  

        st.dataframe(  
            pd.DataFrame(  
                archive_rows  
            ),  
            use_container_width=True,  
            hide_index=True,  
        )  

# --------------------------------------------------------  
# EXPORT  
# --------------------------------------------------------  

st.divider()  

export_bytes = create_planner_excel(  
    planner_data,  
    archived_loads,  
)  

st.download_button(  
    "⬇️ Export Planner to Excel",  
    data=export_bytes,  
    file_name=(  
        f"Campeys_Load_Planner_"  
        f"{current_monday.strftime('%Y%m%d')}.xlsx"  
    ),  
    mime=(  
        "application/vnd.openxmlformats-officedocument."  
        "spreadsheetml.sheet"  
    ),  
    use_container_width=True,  
)

============================================================

HISTORY PAGE

============================================================

def show_history():
if st.button(
"← Back to Home",
key="history_back",
):
go_to("home")
st.rerun()

st.title("📊 SKU History")  

history = load_history()  

if history.empty:  
    st.info(  
        "No SKU history is currently available."  
    )  
    return  

st.session_state.history_df = history  

st.dataframe(  
    history,  
    use_container_width=True,  
    hide_index=True,  
)  

st.subheader(  
    "Archived Loads"  
)  

saved_files = list_saved_load_files()  

if not saved_files:  
    st.info(  
        "No archived load CSVs found."  
    )  
    return  

selected_load = st.selectbox(  
    "Select archived load",  
    saved_files,  
)  

if selected_load:  
    result = get_github_file(  
        f"{GITHUB_FOLDER}/{selected_load}"  
    )  

    if result:  
        st.download_button(  
            "⬇️ Download Archived CSV",  
            data=result["content"].encode(  
                "utf-8-sig"  
            ),  
            file_name=selected_load,  
            mime="text/csv",  
            use_container_width=True,  
        )

============================================================

HISTORIC CSV IMPORT

============================================================

def normalise_history_date(value):
parsed = pd.to_datetime(
value,
errors="coerce",
dayfirst=True,
)

if pd.isna(parsed):  
    return None  

return parsed.strftime(  
    "%Y-%m-%d %H:%M:%S"  
)

def extract_historic_csv_data(
uploaded_file,
):
try:
content = uploaded_file.getvalue()

df = pd.read_csv(  
        io.BytesIO(content),  
        dtype=str,  
        encoding="utf-8-sig",  
    )  

    df.columns = [  
        str(column).strip()  
        for column in df.columns  
    ]  

    if "Item Code" not in df.columns:  
        return []  

    load_ref = ""  

    if "Load Ref" in df.columns:  
        values = (  
            df["Load Ref"]  
            .dropna()  
            .astype(str)  
            .str.strip()  
        )  

        if not values.empty:  
            load_ref = values.iloc[0]  

    date_value = ""  

    for column in [  
        "Date",  
        "Date Submitted",  
    ]:  
        if column in df.columns:  
            values = (  
                df[column]  
                .dropna()  
                .astype(str)  
                .str.strip()  
            )  

            if not values.empty:  
                date_value = (  
                    normalise_history_date(  
                        values.iloc[0]  
                    )  
                    or ""  
                )  
                break  

    if "Quantity" not in df.columns:  
        return []  

    records = []  

    grouped = (  
        df.groupby("Item Code")["Quantity"]  
        .sum()  
    )  

    for item_code, quantity in grouped.items():  
        try:  
            quantity_value = float(  
                str(quantity)  
                .replace(",", "")  
            )  
        except Exception:  
            continue  

        records.append(  
            {  
                "Date": date_value,  
                "Load Ref": load_ref,  
                "Item Code": str(  
                    item_code  
                ).strip(),  
                "Quantity": quantity_value,  
            }  
        )  

    return records  

except Exception:  
    return []

def import_historic_csvs(
imported_records,
):
history = load_history()

incoming = pd.DataFrame(  
    imported_records  
)  

if incoming.empty:  
    return False, "No records to import."  

if history.empty:  
    combined = incoming  
else:  
    combined = pd.concat(  
        [  
            history,  
            incoming,  
        ],  
        ignore_index=True,  
    )  

combined = combined.drop_duplicates(  
    subset=[  
        "Date",  
        "Load Ref",  
        "Item Code",  
        "Quantity",  
    ]  
)  

return save_history(  
    combined  
)

def show_import():
if st.button(
"← Back to Home",
key="import_back",
):
go_to("home")
st.rerun()

st.title("📥 Historic CSV Import")  

st.write(  
    "Upload historic sender CSVs to rebuild or extend "  
    "the SKU history."  
)  

uploaded_files = st.file_uploader(  
    "Historic CSV files",  
    type=["csv"],  
    accept_multiple_files=True,  
)  

if not uploaded_files:  
    return  

imported_records = []  

for uploaded_file in uploaded_files:  
    records = extract_historic_csv_data(  
        uploaded_file  
    )  

    imported_records.extend(  
        records  
    )  

if not imported_records:  
    st.warning(  
        "No usable records were found."  
    )  
    return  

preview = pd.DataFrame(  
    imported_records  
)  

st.subheader(  
    "Import Preview"  
)  

st.dataframe(  
    preview,  
    use_container_width=True,  
    hide_index=True,  
)  

if st.button(  
    "Import Historic Data",  
    type="primary",  
    use_container_width=True,  
):  
    success, message = (  
        import_historic_csvs(  
            imported_records  
        )  
    )  

    if success:  
        st.success(  
            "Historic data imported successfully."  
        )  
    else:  
        st.error(  
            f"Import failed: {message}"  
        )

============================================================

CONTACTS

============================================================

def show_contacts():
if st.button(
"← Back to Home",
key="contacts_back",
):
go_to("home")
st.rerun()

st.title("📇 Campeys Contacts")  

result = get_github_file(  
    CONTACT_LIST_PATH  
)  

if not result:  
    st.info(  
        "No contact list was found."  
    )  
    return  

st.text(  
    result["content"]  
)

============================================================

HOME

============================================================

def show_home():
st.title(
"📦 KP to Campeys SSCC Sender"
)

st.write(  
    "Select a tool below."  
)  

row1 = st.columns(3)  

with row1[0]:  
    if st.button(  
        "📦 SSCC Sender",  
        key="open_sender",  
        use_container_width=True,  
    ):  
        go_to("sender")  
        st.rerun()  

with row1[1]:  
    if st.button(  
        "📊 SKU History",  
        key="open_history",  
        use_container_width=True,  
    ):  
        go_to("history")  
        st.rerun()  

with row1[2]:  
    if st.button(  
        "🖥️ AutoStore",  
        key="open_autostore",  
        use_container_width=True,  
    ):  
        st.link_button(  
            "Open AutoStore",  
            AUTOSTORE_URL,  
            use_container_width=True,  
        )  

row2 = st.columns(3)  

with row2[0]:  
    if st.button(  
        "📇 Campeys Contacts",  
        key="open_contacts",  
        use_container_width=True,  
    ):  
        go_to("contacts")  
        st.rerun()  

with row2[1]:  
    if st.button(  
        "📋 Open Load Planner",  
        key="open_planner",  
        use_container_width=True,  
    ):  
        go_to("planner")  
        st.rerun()  

with row2[2]:  
    if st.button(  
        "📥 Historic CSV Import",  
        key="open_import",  
        use_container_width=True,  
    ):  
        go_to("import")  
        st.rerun()  

st.divider()  

st.info(  
    "The Load Planner uses your existing Campeys Transport "  
    "Requests workbook. Upload it from the planner to "  
    "populate the collection slots."  
)

============================================================

ROUTING

============================================================

if st.session_state.page == "home":
show_home()

elif st.session_state.page == "sender":
show_sender()

elif st.session_state.page == "history":
show_history()

elif st.session_state.page == "planner":
show_planner()

elif st.session_state.page == "import":
show_import()

elif st.session_state.page == "contacts":
show_contacts()

else:
st.session_state.page = "home"
st.rerun()